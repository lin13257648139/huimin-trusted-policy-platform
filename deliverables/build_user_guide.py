from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_submission_doc import (
    AMBER,
    BLUE,
    GRAY,
    MINT,
    NAVY,
    TEAL,
    add_page_number,
    add_run,
    body,
    bullet,
    callout,
    heading,
    page_break,
    table,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "惠民智导_软件使用说明.docx"


def make_doc() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(GRAY)
    for name, size in (("Title", 27), ("Heading 1", 18), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(header, "惠民智导 · 软件使用说明", color=TEAL, size=8.5)
    add_page_number(section.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(78)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "惠民智导可信政策智能平台", bold=True, color="000000", size=28)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "软件安装 运行 演示与故障处理说明", bold=True, color=NAVY, size=18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "V1.0.0 · Windows 本地部署", color=GRAY, size=11)
    callout(
        doc,
        "用途",
        "本说明用于首次安装、现场演示和运行维护。系统输出是政策导办建议，不替代行政机关审批；模型不可用时会降级到确定性安全底座。",
        BLUE,
    )
    page_break(doc)

    heading(doc, "01 运行条件", 1)
    table(
        doc,
        ["项目", "最低要求", "完整 AI 模式建议"],
        [
            ["操作系统", "Windows 10 或 11，64 位", "Windows 11"],
            ["Python", "3.11 或 3.12", "3.12"],
            ["内存", "8GB", "16GB 及以上"],
            ["显卡", "不要求，可运行确定性模式", "NVIDIA RTX 4060 8GB 或同等级"],
            ["磁盘", "软件与依赖约 2GB", "模型另需约 5GB"],
            ["浏览器", "Edge 或 Chrome 当前版本", "Edge 或 Chrome 当前版本"],
        ],
        [3.3, 6.2, 6.7],
    )
    heading(doc, "1.1 两种运行方式", 2)
    bullet(doc, "确定性模式：无需下载模型，完整运行规则裁决、证据展示、材料流程和治理演示，适合首次验收。")
    bullet(doc, "混合模式：在本机安装 Ollama 与指定模型后，增加自然语言结构化抽取和向量检索。最终资格状态仍由规则引擎裁决。")
    callout(doc, "建议", "第一次启动先使用确定性模式验证环境，再安装模型。这样更容易区分安装问题与模型运行问题。", MINT)
    page_break(doc)

    heading(doc, "02 首次安装与启动", 1)
    heading(doc, "2.1 解压与安装", 2)
    body(doc, "把软件运行包完整解压到不含特殊权限限制的目录，例如 D:\\HuiminGuide。不要直接在 ZIP 压缩包内部运行脚本。打开 PowerShell，进入解压目录后执行：")
    table(
        doc,
        ["顺序", "命令", "作用"],
        [
            ["1", ".\\setup.ps1", "创建虚拟环境并安装锁定依赖"],
            ["2", ".\\start.ps1 -AIMode deterministic", "以确定性安全模式启动"],
            ["3", ".\\status.ps1", "查看端口、进程和健康状态"],
            ["4", ".\\logs.ps1", "查看最近启动日志"],
            ["5", ".\\stop.ps1", "安全停止本项目进程"],
        ],
        [1.6, 6.3, 8.3],
    )
    body(doc, "启动成功后访问 http://127.0.0.1:8000。若8000端口已被占用，执行 .\\start.ps1 -Port 8010，并访问 http://127.0.0.1:8010。")
    callout(doc, "执行策略提示", "若系统阻止本地脚本，可在当前 PowerShell 会话执行 Set-ExecutionPolicy -Scope Process Bypass，然后重新运行脚本。该设置只影响当前窗口。", AMBER)
    page_break(doc)

    heading(doc, "03 启用本地 AI", 1)
    heading(doc, "3.1 准备模型", 2)
    body(doc, "安装 Ollama 后，按模型官方方式下载生成模型 qwen3:4b-instruct 与向量模型 qwen3-embedding:0.6b。模型权重不包含在软件包中。")
    table(
        doc,
        ["命令", "说明"],
        [
            ["ollama list", "确认两个模型均已安装"],
            [".\\start-ai-runtime.ps1", "检查并启动本地模型运行服务"],
            [".\\start.ps1 -AIMode hybrid", "启用结构化抽取与混合检索"],
            [".\\start.ps1 -AIMode auto", "模型齐全时自动选用混合模式，否则安全降级"],
        ],
        [7.0, 9.2],
    )
    heading(doc, "3.2 判断是否启用成功", 2)
    bullet(doc, "启动脚本应显示 AI 模式为 hybrid。")
    bullet(doc, "系统状态页应显示生成模型和向量模型可用。")
    bullet(doc, "若模型不可用，系统不会伪装成功，而会显示降级状态并继续提供确定性导办。")
    page_break(doc)

    heading(doc, "04 功能与演示路径", 1)
    table(
        doc,
        ["端口", "功能", "建议演示"],
        [
            ["居民端", "诉求输入、主动补问、资格预判、证据与材料清单", "载入示例家庭，展示 UNKNOWN 与下一步"],
            ["材料预检", "图片或 PDF 选择反馈、OCR 候选、缺件与质量提示", "分别上传清晰与模糊材料"],
            ["工作人员端", "结构化摘要、缺失事实、风险和政策依据", "说明导办不替代人工审核"],
            ["治理端", "政策版本、候选规则、影响分析和人工发布门", "展示模型不能直接发布规则"],
            ["场景验证", "固定案例、时间边界、紧急风险和注入防护", "按 A 至 F 顺序演示"],
            ["可信度中心", "数据、模型、测试、安全和依赖状态", "说明指标样本量与边界"],
        ],
        [2.7, 6.7, 6.8],
    )
    callout(doc, "材料上传反馈", "选择文件后页面会显示文件名、数量、大小和可清空状态。出现“已成功选择”只表示浏览器已接收文件，不表示材料已经通过业务审核。", MINT)
    page_break(doc)

    heading(doc, "05 管理员令牌与隐私", 1)
    heading(doc, "5.1 联网或正式部署", 2)
    body(doc, "仅在本机演示时可使用 development 环境。联网、局域网共享或正式部署前，必须复制 .env.example 为 .env，并设置 APP_ENV=production 与至少32位随机 ADMIN_API_TOKEN。生产环境未配置令牌时，程序会拒绝启动。")
    bullet(doc, "不要把真实令牌写入截图、PPT、视频、源代码或提交包。")
    bullet(doc, "管理员接口使用 X-Admin-Token 请求头；普通居民页面不应暴露该令牌。")
    bullet(doc, "咨询记录使用独立管理令牌，数据库只保存令牌哈希。")
    heading(doc, "5.2 数据处理", 2)
    bullet(doc, "演示优先使用合成案例，不上传真实身份证、银行卡和未脱敏医疗材料。")
    bullet(doc, "运行数据库、日志、PID、端口记录和模型权重不进入公开软件包。")
    bullet(doc, "停止演示后检查 storage 目录，按授权范围处理现场产生的数据。")
    page_break(doc)

    heading(doc, "06 常见问题", 1)
    table(
        doc,
        ["现象", "原因", "处理"],
        [
            ["8000端口被占用", "已有服务或其他程序监听", "先运行 .\\status.ps1；必要时改用 -Port 8010"],
            ["提示尚未初始化", "虚拟环境不存在", "运行 .\\setup.ps1"],
            ["启动后健康检查失败", "依赖、端口或配置错误", "运行 .\\logs.ps1 查看 storage\\server.err.log"],
            ["页面仍显示旧样式", "浏览器或 Service Worker 缓存", "强制刷新，必要时清除该站点缓存后重开"],
            ["AI状态不可用", "Ollama未启动或模型标签不匹配", "运行 ollama list，再执行 start-ai-runtime.ps1"],
            ["上传后无结果", "只选择了文件但尚未执行检查，或文件超限", "确认已显示文件清单，再点击材料检查并查看提示"],
            ["stop.ps1拒绝停止", "端口上的进程不属于本项目环境", "按日志和PID人工确认，避免误杀其他程序"],
        ],
        [4.1, 5.0, 7.1],
    )
    page_break(doc)

    heading(doc, "07 验收清单", 1)
    table(
        doc,
        ["检查项", "通过标准"],
        [
            ["安装", "setup.ps1无错误，锁定依赖安装完成"],
            ["启动", "健康检查通过，浏览器可访问，页面版本为V1.0.0"],
            ["路由", "居民端、工作人员端、治理端和场景验证页面可切换"],
            ["上传", "选择文件后显示文件名、数量、大小和清空按钮"],
            ["证据", "导办结果能查看逐字政策原文、来源和适用日期"],
            ["安全", "信息不足返回UNKNOWN；注入和越权请求被阻断"],
            ["降级", "关闭Ollama后仍能在确定性模式运行，并显示真实状态"],
            ["停止", "stop.ps1只停止本项目进程，端口释放"],
        ],
        [5.1, 11.1],
    )
    callout(doc, "边界", "真实群众可用性、办事成功率、跑腿减少比例和基层实际采用情况，必须由真实试点产生。本软件包中的内部评测不能替代这些外部证据。", AMBER)

    doc.core_properties.title = "惠民智导可信政策智能平台软件使用说明"
    doc.core_properties.subject = "安装运行演示与故障处理"
    doc.core_properties.author = "策马奔腾"
    doc.core_properties.keywords = "惠民智导, 安装, 运行, 演示, 故障处理, 隐私"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make_doc()
