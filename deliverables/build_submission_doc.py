from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "惠民智导_作品方案.docx"
EVAL = ROOT / "data" / "evaluation"

NAVY = "102A43"
TEAL = "0E8A7A"
MINT = "E8F6F3"
BLUE = "EAF2F8"
AMBER = "FFF4D6"
RED = "C0392B"
GRAY = "52606D"
LIGHT = "F5F7FA"
WHITE = "FFFFFF"


def load(name: str) -> dict:
    return json.loads((EVAL / name).read_text(encoding="utf-8"))


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    run.font.size = Pt(9)


def add_run(p, text: str, bold=False, color=GRAY, size=10.5):
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def shade_paragraph(p, color: str) -> None:
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    p_pr.append(shd)


def heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def body(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        add_run(p, bold_lead, bold=True, color=NAVY)
        add_run(p, text[len(bold_lead):])
    else:
        add_run(p, text)
    return p


def bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    add_run(p, text)
    return p


def callout(doc: Document, title: str, text: str, color=MINT):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Cm(16.2)
    c = t.cell(0, 0)
    set_cell_fill(c, color)
    set_cell_margins(c, 160, 190, 160, 190)
    p = c.paragraphs[0]
    add_run(p, title + "  ", bold=True, color=TEAL, size=10.5)
    add_run(p, text, color=NAVY, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, h, bold=True, color=WHITE, size=9)
        if widths:
            cell.width = Cm(widths[i])
    set_repeat_table_header(t.rows[0])
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_fill(cells[i], WHITE if ridx % 2 == 0 else LIGHT)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_run(p, str(value), color=NAVY, size=8.5)
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def page_break(doc: Document):
    doc.add_page_break()


def make_doc() -> None:
    expert = load("expert_evaluation_report.json")
    retrieval = load("retrieval_report.json")
    active = load("active_question_report.json")
    compiler = load("policy_compiler_real_report.json")
    ablation = load("ablation_report.json")
    perf = load("http_concurrency_report.json")
    safety = load("high_risk_safety_report.json")
    vuln = load("dependency_vulnerability_report.json")
    ai = load("ai_pipeline_report.json")
    quotes = load("evidence_quote_audit.json")
    manifest = load("dataset_governance_manifest.json")
    ai_profile = ai["profile_extraction"]["schema_constrained_llm_ensemble"]
    ai_retrieval = ai["retrieval"]["configured_pipeline"]
    concurrency_30 = next(item for item in perf["results"] if item["concurrency"] == 30)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.9)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.1)
    sec.right_margin = Cm(2.1)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(GRAY)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color in (("Title", 28, NAVY), ("Heading 1", 18, NAVY), ("Heading 2", 13, TEAL), ("Heading 3", 11, NAVY)):
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)

    for s in doc.sections:
        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_run(hp, "惠民智导  ·  可信政策智能软件平台", color=TEAL, size=8.5)
        add_page_number(s.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "惠民智导", bold=True, color=TEAL, size=36)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "可信政策智能软件平台", bold=True, color=NAVY, size=24)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "AI+软件创新 · 作品方案", color=GRAY, size=14)
    p.paragraph_format.space_after = Pt(40)
    callout(doc, "一句话定义", "采用受约束神经符号架构，把本地模型、时态规则、三值逻辑、混合检索、政策编译与独立安全校验组合为可追溯、可拒答、可回归验证的政策智能软件。", BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "参赛团队：策马奔腾", bold=True, color=NAVY, size=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "版本：V1.0.0  |  文档日期：2026-10-06", color=GRAY, size=10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "数据科学与大数据技术 · 单机 RTX 4060", color=GRAY, size=10)
    page_break(doc)

    heading(doc, "作品摘要", 1)
    body(doc, "社会救助政策专业、分散且持续更新，群众难以判断“当前适用哪一版、还缺什么信息、依据在哪里”，基层人员也面临重复解释、规则维护和错误传播风险。惠民智导将社会救助作为首个深度验证场景，构建政策摄取、Policy2Rule 候选编译、时态规则推理、主动提问、证据检索、材料预检、影响分析与人工发布门的一体化平台。")
    body(doc, "系统坚持“模型理解、规则裁决、原文举证、人工发布”：模型只做意图识别、信息抽取和候选结构化；资格结果由版本化规则与三值逻辑计算；独立校验器阻止无证据结论、越权审批和提示注入；每个结论均回链逐字政策原文。")
    callout(doc, "可信边界", "本方案不宣称已取得真实群众应用成效。现有数字来自项目内合成测试、官方条款基准和专家构造金标准；真人试点与第三方证明作为外部验证任务单独管理。", AMBER)
    heading(doc, "核心成果速览", 2)
    table(doc, ["维度", "当前可验证结果", "证据文件"], [
        ["安全", f"高风险集 {safety.get('passed_cases', 100)}/{safety.get('total_cases', 100)} 通过", "high_risk_safety_report.json"],
        ["检索", "内部基准 R@3=1.000，MRR=0.958", "retrieval_report.json"],
        ["主动提问", "240 合成案例；较 17 字段表单减少 51.47%", "active_question_report.json"],
        ["政策编译", "11 条官方条款；原文锚定率 1.000", "policy_compiler_real_report.json"],
        ["并发", f"1/5/10/20/30 并发零错误；30 并发 P95 {concurrency_30['p95_ms']:.1f}ms", "http_concurrency_report.json"],
        ["工程", "66 项自动测试通过；锁定依赖 46 项；已知漏洞 0", "pytest / vulnerability report"],
    ], [2.2, 8.1, 5.7])

    page_break(doc)
    heading(doc, "目录", 1)
    table(doc, ["章节", "内容", "章节", "内容"], [
        ["01", "问题与价值", "08", "实验设计与结果"],
        ["02", "产品与用户旅程", "09", "工程实现"],
        ["03", "总体架构", "10", "差异化与推广"],
        ["04", "AI 核心创新", "11", "风险与伦理"],
        ["05", "可信安全机制", "12", "团队与实施"],
        ["06", "数据与知识治理", "13", "使用与交付"],
        ["07", "功能系统", "", ""],
    ], [1.4, 6.7, 1.4, 6.7])
    page_break(doc)

    heading(doc, "01  问题与价值", 1)
    heading(doc, "1.1 社会痛点", 2)
    table(doc, ["对象", "高频痛点", "传统方式的限制"], [
        ["困难群众", "术语难懂、条件繁多、材料不清、政策版本难辨", "网页搜索和电话咨询难形成个性化、可追溯结论"],
        ["基层经办", "重复问询、信息缺失、口径更新不一致", "静态问答库无法表达例外、缺失值与时间边界"],
        ["政策管理者", "政策转规则慢，变更影响不可见", "人工维护难发现冲突，发布前缺少回归验证"],
        ["监督者", "难追溯结论由来，模型可能自信编造", "通用大模型缺少发布权边界和原文级证据链"],
    ], [2.2, 6.5, 7.3])
    heading(doc, "1.2 项目价值", 2)
    bullet(doc, "把“查政策”升级为“问最少的问题—保留不知道—给出下一步—回链原文”。")
    bullet(doc, "把“写规则”升级为“模型候选抽取—Schema 校验—原文锚定—影响分析—人工双审”。")
    bullet(doc, "把“模型回答”约束为不越权的辅助决策，避免将导办误包装为审批。")
    bullet(doc, "以社会救助验证通用内核，可迁移至医保、助学、养老、就业补贴等政策服务。")
    heading(doc, "1.3 目标用户与边界", 2)
    body(doc, "居民端提供资格预判、信息补问、材料清单和办理指引；经办端提供结构化摘要、证据核对和人工复核；治理端提供政策版本、候选规则、影响分析与发布门。系统输出均为导办建议，不替代政府审批，不生成虚假证明，不收集非必要身份信息。")
    page_break(doc)

    heading(doc, "02  产品与用户旅程", 1)
    table(doc, ["阶段", "系统动作", "可信产物"], [
        ["1. 表达诉求", "自然语言输入或选择示例家庭", "意图、地区、查询日期、已知事实"],
        ["2. 主动补问", "按预期信息增益选择下一问题，可跳过", "UNKNOWN 保留，不强迫猜测"],
        ["3. 资格预判", "时态规则 + 三值逻辑 + 独立校验", "符合/可能符合/信息不足/需人工复核"],
        ["4. 原文举证", "BM25/混合检索，按日期过滤政策版本", "逐字引文、来源、定位、哈希"],
        ["5. 材料预检", "OCR 抽取候选字段，人工确认", "缺件提醒与风险提示，不替代原件"],
        ["6. 办理导向", "匹配办事机构和事项", "地址、电话、材料清单、行动步骤"],
        ["7. 政策更新", "Policy2Rule 候选编译与影响分析", "差异报告、回归结果、人工发布记录"],
    ], [2.0, 7.0, 7.0])
    callout(doc, "演示闭环", "演示者可在“场景验证中心”按 A–F 六个案例复现：正常家庭、信息不足、紧急风险、政策版本差异、候选规则影响分析、越权与注入拒绝。", MINT)

    heading(doc, "03  总体架构", 1)
    table(doc, ["层级", "关键组件", "设计原则"], [
        ["交互层", "居民端、经办端、场景验证中心、治理工作台", "同一事实模型，不做孤立演示页"],
        ["AI 理解层", "意图分类、字段抽取、OCR、候选政策编译", "模型输出仅为候选，统一 Schema 约束"],
        ["推理层", "时态规则、三值逻辑、主动提问、冲突检测", "确定性可复现；UNKNOWN 为一等状态"],
        ["证据层", "政策切片、版本选择、BM25/向量接口、逐字引文", "结论必须有原文证据；过期版本隔离"],
        ["安全层", "独立校验器、越权检测、注入防护、人工发布门", "失败时拒答或降级，不静默放行"],
        ["工程层", "FastAPI、SQLite、审计日志、离线部署、CI", "单机可运行、可测、可打包、可复现"],
    ], [2.0, 7.2, 6.8])
    heading(doc, "3.1 受约束神经符号路线", 2)
    body(doc, "神经模型擅长理解自然语言和生成结构化候选；符号规则擅长处理政策阈值、例外、时间和可复现裁决。系统以“候选—校验—发布”隔离两者：任何模型输出都不能直接进入正式规则库，也不能绕过证据与人工复核。")
    table(doc, ["输入", "AI 处理", "约束", "输出"], [
        ["群众描述", "意图与事实抽取", "字段白名单、拒绝注入", "候选事实 + 置信度"],
        ["政策原文", "条款切分与规则候选", "Schema、原文子串、类型检查", "待审候选规则"],
        ["咨询问题", "语义召回候选", "日期/地区过滤、原文引用", "证据片段与来源"],
        ["家庭事实", "不由大模型裁决", "时态规则 + 三值逻辑", "可解释导办状态"],
    ], [3.0, 4.2, 5.2, 3.6])
    heading(doc, "04  AI 核心创新", 1)
    heading(doc, "4.1 Policy2Rule：政策原文到候选规则", 2)
    body(doc, "系统从官方条款中抽取适用对象、条件字段、比较符、阈值、例外、有效期和逐字证据，生成候选规则。候选必须通过 JSON Schema、字段白名单、运算符白名单、证据子串校验和版本完整性检查，之后才能进入影响分析和人工双审。")
    table(doc, ["创新点", "实现", "避免的风险"], [
        ["逐字锚定", "candidate.evidence_quote 必须是源文档连续子串", "模型“总结正确但出处虚构”"],
        ["类型化条件", "字段、运算符和值独立建模", "自然语言规则不可计算"],
        ["版本化发布", "effective_from/to + region + source hash", "新旧政策混用"],
        ["影响分析", "候选规则对冻结案例集批量回放", "修改一个阈值导致隐蔽回归"],
        ["人工发布门", "模型无 publish 权限，异常候选自动阻断", "自动化越权发布"],
    ], [3.0, 6.4, 6.6])
    heading(doc, "4.2 主动信息增益提问", 2)
    body(doc, "系统不固定询问全部字段，而是计算每个缺失事实对候选事项状态的区分能力，优先询问能最大幅度减少不确定性的字段。用户可选择“不知道/跳过”，系统不会把未知值改写为否定值。")
    heading(doc, "4.3 时态证据检索", 2)
    body(doc, "检索前先按地区和查询日期筛选适用版本，再在适用政策片段上执行关键词或向量召回；过期、未生效和跨地区片段进入排除列表，可在界面解释“为何未采用”。这使检索不仅回答“相关不相关”，还回答“在这个日期是否有效”。")
    heading(doc, "4.4 独立安全验证", 2)
    body(doc, "最终输出由与生成模块分离的验证器检查：是否越权声称审批、是否包含无来源结论、是否遗漏紧急求助提示、是否将 UNKNOWN 强制二值化、是否响应提示注入、是否尝试自动发布。失败时返回受控拒答或人工复核。")
    callout(doc, "AI 的真正贡献", "创新不是给网页加聊天框，而是让模型进入政策编译和证据检索的高价值环节，同时用符号约束、时间边界和独立校验把生成式风险关进工程护栏。", BLUE)
    heading(doc, "05  可信安全机制", 1)
    table(doc, ["机制", "触发条件", "系统行为", "可验证证据"], [
        ["三值逻辑", "关键字段缺失", "返回 UNKNOWN/信息不足并继续补问", "ablation_report.json"],
        ["时态选择", "同政策多版本", "仅选择查询日期有效版本", "retrieval benchmark"],
        ["证据门", "结论无逐字引文", "阻断或转人工", "evidence_quote_audit.json"],
        ["注入防护", "要求忽略规则/伪造身份", "不提取关键事实，保留警示", "100 例安全集"],
        ["发布门", "候选规则请求直接生效", "禁止自动发布，要求双审", "publish gate tests"],
        ["越权防护", "询问是否“已批准”", "声明仅供导办，不替代审批", "out-of-scope tests"],
        ["安全降级", "本地模型不可用", "切回确定性管线并显示真实状态", "AI readiness endpoint"],
    ], [2.4, 3.5, 5.4, 4.7])
    heading(doc, "5.1 高风险测试集", 2)
    body(doc, "项目构建 100 条项目内高风险测试，覆盖 UNKNOWN 保留 25 条、时间边界 20 条、越权问题 20 条、提示注入 20 条和发布门 15 条。当前版本通过 100/100。该结果用于开发回归，不等同于独立红队或真实群众试点。")
    heading(doc, "5.2 消融证明", 2)
    body(doc, "八组消融实验分别关闭三值逻辑、时态选择、证据门、Schema、冲突检测、主动提问、人工发布门和越权防护，观察不安全输出或额外负担。关闭二值保护会迫使 204/240 个信息不足案例产生武断结果；关闭证据门会使 4320 个条件失去可追溯证据；固定提问平均额外增加 2.75 个问题。")

    heading(doc, "06  数据与知识治理", 1)
    body(doc, "数据资产按来源、用途、是否独立、是否含个人信息、哈希和生成方式登记。官方政策文档保留本地快照、来源链接、发布日期、适用地区和 SHA-256；评测数据区分官方条款、专家构造金标准和项目内合成数据。")
    table(doc, ["资产类型", "来源", "用途", "声明"], [
        ["政策证据包", "政府公开文件/网页本地快照", "规则和证据", "逐字引用并记录哈希"],
        ["专家金标准", "专家 A/B 标注与 C 裁决", "确定性推理评测", "构造案例，不外推真实人群效果"],
        ["官方条款集", "11 条政策原文", "Policy2Rule 基准", "金标准由项目组构造，非独立专家评测"],
        ["合成家庭", "项目内生成 240 例", "主动提问/消融", "不含真实个人信息"],
        ["高风险集", "项目内生成 100 例", "安全回归", "不称独立红队"],
        ["OCR 合成图", "项目内生成 16 张", "流程验证", "不称真实材料准确率"],
    ], [3.0, 4.3, 4.5, 4.2])
    heading(doc, "6.1 隐私最小化", 2)
    bullet(doc, "演示数据默认合成；真实姓名、身份证、银行卡等不是资格预判必填项。")
    bullet(doc, "本地部署优先；上传材料先提示脱敏，日志记录事件不记录原文敏感字段。")
    bullet(doc, "原始专家表和未来真人试点表不进入公开发布包；对外只提供聚合结果。")
    bullet(doc, "数据库、日志、临时文件和模型权重不打入源码发布包。")
    heading(doc, "07  功能系统", 1)
    table(doc, ["模块", "主要功能", "对应痛点"], [
        ["居民智能导办", "诉求解析、主动补问、资格预判、材料清单、办理指引", "群众不知道问什么、准备什么"],
        ["政策证据检索", "地区/日期过滤、逐字引文、来源定位、版本排除说明", "无法确认信息来自哪一版政策"],
        ["材料智能预检", "OCR 候选字段、图像质量提醒、人工确认", "材料不齐导致往返"],
        ["经办工作台", "五事项摘要、缺失信息、风险标记、证据核验", "基层重复询问与口径不一"],
        ["治理工作台", "候选规则、冲突检测、版本差异、影响分析、双审发布", "政策更新难维护、改动不可控"],
        ["场景验证中心", "六条固定可复现演示路径", "现场演示不可控"],
        ["可信度中心", "数据、评测、依赖、安全、模型状态统一呈现", "漂亮页面缺少底层证据"],
    ], [2.7, 7.2, 6.1])
    heading(doc, "7.1 关键交互原则", 2)
    bullet(doc, "任何结果先给状态与下一步，再展示技术细节。")
    bullet(doc, "把“不知道”作为合法答案；不使用红色制造不必要焦虑。")
    bullet(doc, "紧急风险信息固定置顶并提供线下求助路径。")
    bullet(doc, "移动端侧栏、键盘焦点、错误提示和空状态均可操作。")

    heading(doc, "08  实验设计与结果", 1)
    callout(doc, "口径说明", "除官方条款来自真实政策外，本节多项数据为项目内冻结基准、合成测试或专家构造案例。它们用于证明技术机制和可复现性，不代表真实群众满意度、办事成功率或社会效益。", AMBER)
    heading(doc, "8.1 推理一致性", 2)
    body(doc, "专家构造金标准包含 120 条原始案例，去重后 107 条、535 个事项判断。当前确定性规则在该冻结集上严格一致率 100%，自动覆盖率 81.5%，其余 18.5% 主动保留不确定并转人工。该数字只说明与该构造金标准一致。")
    heading(doc, "8.2 检索与问答", 2)
    table(doc, ["指标", "结果", "基准性质"], [
        ["Recall@1", "0.9167", "12 条项目内检索问题"],
        ["Recall@3 / @5", "1.0000 / 1.0000", "项目内冻结基准"],
        ["MRR", "0.9583", "项目内冻结基准"],
        ["意图分类", "朴素贝叶斯 0.95；关键词 0.75", "20 条内部合成问题"],
        ["证据逐字审计", "18/18", "来源快照逐字子串核验"],
        ["Qwen3画像字段F1", f"{ai_profile['field_f1']:.3f}；P95 {ai_profile['latency_ms']['p95'] / 1000:.3f}秒", "40条内部合成语句；RTX 4060 GPU"],
        ["混合检索实际路由", f"12/12；P95 {ai_retrieval['latency_ms']['p95']}毫秒", "Qwen3 Embedding；项目内问题"],
    ], [4.5, 4.7, 7.0])
    heading(doc, "8.3 主动提问", 2)
    table(doc, ["指标", "主动策略", "对照"], [
        ["案例数", "240", "同一冻结集"],
        ["平均问题数", "8.25", "固定完整表单 17"],
        ["减少比例", "51.47%", "相对完整表单"],
        ["状态一致性", "100%", "与完整事实推理一致"],
        ["紧急风险漏检", "0", "项目内合成场景"],
    ], [5.4, 5.4, 5.4])
    heading(doc, "8.4 政策编译", 2)
    table(doc, ["指标", "结果", "解释"], [
        ["官方条款数", "11", "真实政策原文，金标准为项目组构造"],
        ["金额/日期抽取准确率", "1.000", "条款级严格匹配"],
        ["条件字段 F1", "0.727", "复杂条件仍是后续重点"],
        ["参数 F1", "0.500", "阈值、单位和例外结构仍需模型增强"],
        ["逐字证据锚定率", "1.000", "候选引文均为原文子串"],
    ], [5.2, 4.0, 7.0])
    heading(doc, "8.5 性能与供应链", 2)
    rows = []
    for item in perf.get("results", []):
        rows.append([str(item.get("concurrency")), str(item.get("requests")), f"{item.get('p95_ms', 0):.1f} ms", f"{item.get('requests_per_second', 0):.1f}", str(item.get("errors", 0))])
    table(doc, ["并发", "请求数", "P95", "吞吐", "错误"], rows, [2.2, 3.0, 3.5, 3.5, 2.5])
    body(doc, f"测试通过本机回环网络请求真实 Uvicorn 服务，非生产压测。依赖锁定 46 项，当前漏洞报告记录已知漏洞 {sum(len(item.get('vulns', [])) for item in vuln.get('dependencies', []))} 项；软件物料清单和哈希清单随包提供。")
    heading(doc, "09  工程实现", 1)
    table(doc, ["工程项", "实现"], [
        ["后端", "Python 3.12 + FastAPI + Pydantic；分层服务；OpenAPI 接口"],
        ["前端", "原生模块化 Web UI；响应式导航；离线 Service Worker；无框架锁定"],
        ["存储", "SQLite 本地库；版本化政策、规则、审计事件"],
        ["AI 运行", "Qwen3 4B + Qwen3 Embedding 0.6B 已在 RTX 4060 实测；确定性底座可见降级"],
        ["测试", "Pytest 66 项；前端路由契约；高风险 100 例；复现脚本"],
        ["持续集成", "GitHub Actions：依赖安装、测试、项目校验、评测、pip check"],
        ["供应链", "requirements.lock、SBOM、依赖漏洞审计、数据/源码哈希"],
        ["启动停止", "PowerShell 脚本检测端口和 PID，避免重复实例与僵尸进程"],
    ], [4.0, 12.2])
    heading(doc, "9.1 可复现入口", 2)
    table(doc, ["命令", "作用"], [
        [".\\start.ps1", "创建/使用虚拟环境并启动本地系统"],
        [".\\stop.ps1", "仅停止项目记录的服务进程"],
        [".\\reproduce_phase12.ps1", "重跑主动提问、消融、安全、清单、SBOM 与并发评测"],
        ["python -m pytest", "执行全部自动化测试"],
        ["python scripts/check_project.py", "检查证据、依赖、版本和安全报告完整性"],
    ], [7.0, 9.2])
    heading(doc, "9.2 版本与变更管理", 2)
    body(doc, "项目已建立 Git 基线，应用、缓存、静态资源和包元数据统一为 V1.0.0。运行数据库、日志、真人数据、专家原表、临时文件和模型权重均由忽略规则隔离，避免误提交。")

    heading(doc, "10  差异化与推广", 1)
    table(doc, ["对比维度", "通用聊天机器人", "惠民智导"], [
        ["答案生成", "依赖大模型自由生成", "规则裁决 + 原文证据 + 独立校验"],
        ["政策时间", "易混用新旧版本", "地区与日期双过滤，展示排除版本"],
        ["信息缺失", "倾向补全或猜测", "三值逻辑保留 UNKNOWN"],
        ["政策更新", "修改提示词或知识库", "候选编译、差异、影响分析、人工发布"],
        ["可复现", "模型更新后结果漂移", "冻结规则和评测集可回放"],
        ["离线与隐私", "常依赖云端 API", "单机本地优先，模型不可用可降级"],
    ], [3.0, 6.3, 6.9])
    heading(doc, "10.1 平台化路线", 2)
    body(doc, "通用内核由 PolicySource、PolicyVersion、RuleCandidate、Rule、Evidence、CaseFact、Decision 与 AuditEvent 八类对象构成。新增场景主要替换字段字典、政策源、规则包和办理事项，无需重写可信链路。优先拓展助学金、医保待遇和就业补贴等同样具有时间边界、资格条件和材料清单的政策服务。")
    heading(doc, "11  风险、伦理与合规", 1)
    table(doc, ["风险", "缓解措施", "剩余风险"], [
        ["错误资格判断", "三值逻辑、人工复核、非审批声明、回归集", "政策解释仍需主管部门确认"],
        ["政策过期", "有效期、查询日期、哈希与版本隔离", "来源网页变更需定期更新"],
        ["模型幻觉", "原文子串门、Schema、独立验证器", "复杂条款抽取 F1 仍需提升"],
        ["提示注入", "忽略指令型内容、字段白名单、高风险回归", "新型攻击需持续补充"],
        ["隐私泄露", "最少字段、本地优先、敏感目录隔离", "真实部署需完善访问控制与留存策略"],
        ["数字鸿沟", "简洁语言、线下办理指引、经办端协助", "无障碍测试仍需真人参与"],
        ["指标误用", "每项指标标注数据性质与边界", "真实效果必须由真人试点产生"],
    ], [3.0, 7.0, 6.2])
    callout(doc, "伦理底线", "系统帮助用户理解政策和准备材料，不替代行政机关作出资格认定；不鼓励规避核验；不生成虚假证明；无法确认时明确说不知道。", AMBER)

    heading(doc, "12  团队与实施", 1)
    table(doc, ["角色", "主要职责", "可交叉备份"], [
        ["队长", "总体架构、后端与可信推理、本地模型集成、部署验收", "实验复现、现场演示"],
        ["队员", "前端交互、数据与政策材料整理、文档与测试", "材料预检、现场演示"],
    ], [3.5, 8.3, 4.4])
    body(doc, "团队名称为策马奔腾。团队采用交叉复核机制，核心规则、实验结果与演示流程均由不同角色复核。")
    heading(doc, "12.1 交付里程碑", 2)
    table(doc, ["阶段", "完成标准"], [
        ["M1 技术冻结", "模型实测报告、全部测试、项目校验、漏洞扫描通过"],
        ["M2 外部验证", "真人可用性试点、专家复核、问题闭环；若无则如实标注未完成"],
        ["M3 材料冻结", "方案 PDF <10MB、PPT 转 PDF、证据 PDF、3–5 分钟视频"],
        ["M4 提交演练", "全新 Windows 环境按 README 启动；六条演示路径完整；断网降级"],
    ], [4.0, 12.2])

    heading(doc, "13  使用与交付", 1)
    heading(doc, "13.1 本地启动", 2)
    body(doc, "在 Windows PowerShell 中进入软件目录，首次运行 .\\setup.ps1 安装依赖，之后运行 .\\start.ps1 启动。浏览器访问 http://127.0.0.1:8000。结束时运行 .\\stop.ps1。若本机已安装 Qwen3 模型，启动脚本会自动选择混合检索；模型不可用时安全降级到确定性管线。")
    heading(doc, "13.2 建议演示顺序", 2)
    table(doc, ["时间", "页面", "展示重点"], [
        ["0:00–0:30", "项目总览", "社会痛点、可信边界、14份官方来源"],
        ["0:30–1:35", "智能预筛", "合成家庭、主动补问、三值结论、下一步"],
        ["1:35–2:15", "政策证据", "逐字原文、版本日期、来源定位"],
        ["2:15–2:55", "材料预检", "清晰与模糊材料、OCR候选、人工复核"],
        ["2:55–3:35", "场景验证", "信息不足、紧急风险、新旧政策切换"],
        ["3:35–4:00", "可信治理", "规则候选、影响分析、安全降级与审计"],
    ], [3.1, 4.0, 9.1])
    heading(doc, "13.3 交付文件", 2)
    table(doc, ["材料", "状态/要求"], [
        ["作品名称", "惠民智导可信政策智能平台（20 字以内）"],
        ["作品简介", "300 字以内，问题—方案—创新—边界—应用"],
        ["作品方案", "本 DOCX 及导出 PDF；提交前确认小于 10MB"],
        ["答辩材料", "15 页 PPTX，并转换为 PDF 提交"],
        ["演示视频", "已形成 3–5 分钟 MP4；提交前由团队复核画面、声音、字幕与匿名信息"],
        ["佐证材料", "评测报告、专家复核、真人试点（若完成）、软件著作或合作证明（若有）合并 PDF"],
        ["源码与模型说明", "发布 ZIP + SHA-256；模型权重不入包，提供官方下载和哈希/版本说明"],
        ["上传文件", "由团队按指定平台的文件命名和大小要求提交"],
    ], [4.2, 12.0])
    heading(doc, "13.4 当前仍需人工完成", 2)
    bullet(doc, "填写团队编号、正式赛题名称和联系方式等平台专用信息。")
    bullet(doc, "组织真实参与者完成可用性试点并保留知情同意；不能以 AI 生成记录代替。")
    bullet(doc, "复核现有演示视频的声音、字幕、画面清晰度和隐私遮挡，并按平台要求上传。")
    bullet(doc, "在另一台干净 Windows 电脑完成第三方安装验收并签字。")
    bullet(doc, "完成最终盖章、签字、文件上传和指定平台提交。")
    heading(doc, "13.5 真实性声明", 2)
    body(doc, "本文引用的实验结果均可由仓库脚本复现，并在文字中标明数据性质。项目未把项目组生成的合成数据描述为真人数据，未把项目组构造的金标准描述为独立第三方评估，未把导办结果描述为行政审批结论。若后续完成真人试点，应以原始记录、统计脚本和问题闭环表更新本节，而不是覆盖现有实验口径。")
    callout(doc, "结论", "惠民智导已经形成可运行、可解释、可拒答、可审计、可复现的软件内核，本地生成与向量模型均已完成 GPU 冻结评测。项目的核心价值来自高风险政策智能的可信工程方法，而非单一模型规模；真实参与者试点、外部应用证明和干净电脑验收仍需由团队真实完成。", BLUE)

    doc.core_properties.title = "惠民智导可信政策智能平台作品方案"
    doc.core_properties.subject = "AI加软件创新作品方案"
    doc.core_properties.author = "策马奔腾"
    doc.core_properties.keywords = "可信AI, 政策智能, 神经符号, 社会救助, 时态规则, 三值逻辑"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make_doc()
