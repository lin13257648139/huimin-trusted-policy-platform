from __future__ import annotations

import json
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "最终交付"
SHOTS = OUT / "截图证据"
PREFIX = "AIC-2026-71323388-AI+软件创新赛-惠民智导可信政策智能平台"
PROPOSAL = OUT / f"{PREFIX}-作品方案.docx"
MANUAL = OUT / f"{PREFIX}-软件使用说明.docx"

GREEN = "075F50"
GREEN_DARK = "12352F"
GREEN_PALE = "EAF5F1"
BLUE_PALE = "EDF4FB"
GRAY_PALE = "F4F6F5"
GRAY_LINE = "D9DFDC"
AMBER_PALE = "FFF6E4"
BLACK = "000000"


def load_evaluation_metrics() -> dict[str, object]:
    evaluation = ROOT / "data" / "evaluation"
    coverage = json.loads((evaluation / "coverage_report.json").read_text(encoding="utf-8"))
    policy = json.loads(
        (evaluation / "policy_compiler_real_report.json").read_text(encoding="utf-8")
    )
    ai_pipeline = json.loads(
        (evaluation / "ai_pipeline_report.json").read_text(encoding="utf-8")
    )
    active_questions = json.loads(
        (evaluation / "active_question_report.json").read_text(encoding="utf-8")
    )
    pytest_root = ET.parse(evaluation / "pytest_junit.xml").getroot()
    pytest_suite = pytest_root.find(".//testsuite")
    if pytest_suite is None:
        raise RuntimeError("pytest_junit.xml缺少testsuite")
    pytest_metrics = {key: int(pytest_suite.attrib.get(key, "0")) for key in ("tests", "failures", "errors", "skipped")}
    if pytest_metrics["failures"] or pytest_metrics["errors"]:
        raise RuntimeError(f"Python测试未全通过：{pytest_metrics}")
    browser = json.loads(
        (evaluation / "browser_e2e_report.json").read_text(encoding="utf-8")
    )["summary"]
    if browser["unexpected"] or browser["flaky"]:
        raise RuntimeError(f"浏览器测试未稳定通过：{browser}")
    return {
        "coverage": coverage["totals"],
        "policy": policy,
        "ai_pipeline": ai_pipeline,
        "active_questions": active_questions,
        "pytest": pytest_metrics,
        "browser": browser,
    }


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = GRAY_LINE, size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def set_run_font(run, east_asia: str, latin: str | None = None, size: float | None = None,
                 bold: bool | None = None, color: str = BLACK) -> None:
    run.font.name = latin or east_asia
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def configure_document(doc: Document, short_title: str) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("26332F")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(5)

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(14)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 18, 18, 9),
        ("Heading 2", 14, 14, 7),
        ("Heading 3", 11.5, 10, 4),
    ):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = short_title
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run_font(header.runs[0], "微软雅黑", size=8.5, color="68736F")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("策马奔腾队  |  惠民智导 v2.0.0  |  2026 AI 软件创新赛")
        set_run_font(run, "微软雅黑", size=8.5, color="68736F")


def cover(doc: Document, doc_type: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    run = p.add_run("2026 AI 软件创新赛")
    set_run_font(run, "微软雅黑", size=14, bold=True, color=GREEN)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("惠民智导可信政策智能平台")
    ppr = p._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(doc_type)
    set_run_font(run, "微软雅黑", size=20, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    run = p.add_run(subtitle)
    set_run_font(run, "微软雅黑", size=11, color="56645F")

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(4.0)
    table.columns[1].width = Cm(8.0)
    items = (
        ("参赛团队", "策马奔腾队"),
        ("参赛编号", "AIC-2026-71323388"),
        ("作品版本", "v2.0.0"),
        ("材料日期", "2026年10月21日"),
    )
    for row, (label, value) in zip(table.rows, items, strict=False):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], GREEN_PALE)
        for index, cell in enumerate(row.cells):
            set_cell_margins(cell, 145, 160, 145, 160)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run_font(run, "微软雅黑", size=10, bold=index == 0)
    set_table_borders(table)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def paragraph(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, "宋体", size=10.5, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, "宋体", size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, "宋体", size=10.5)


def bullets(doc: Document, items: list[str] | tuple[str, ...], numbered: bool = False) -> None:
    style = "Normal" if numbered else "List Bullet"
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.add_run(f"{index}.  {item}" if numbered else item)
        for run in p.runs:
            set_run_font(run, "宋体", size=10.3)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for idx, width in enumerate(widths):
            table.columns[idx].width = Cm(width)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        cell.text = text
        set_cell_shading(cell, GREEN_DARK)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, "微软雅黑", size=9.2, bold=True, color="FFFFFF")
        set_cell_margins(cell, 130, 110, 130, 110)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.text = str(value)
            if row_index % 2:
                set_cell_shading(cell, GRAY_PALE)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    set_run_font(run, "宋体", size=8.8)
            set_cell_margins(cell, 115, 105, 115, 105)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_picture(doc: Document, filename: str, caption: str, width_cm: float = 16.0) -> None:
    path = SHOTS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    run = c.add_run(caption)
    set_run_font(run, "微软雅黑", size=8.8, color="596762")


def add_section_title(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, "微软雅黑", size={1: 18, 2: 14, 3: 11.5}[level], bold=True)


def page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc: Document, entries: list[str]) -> None:
    add_section_title(doc, "目录", 1)
    for index, text in enumerate(entries, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"{index:02d}  {text}")
        set_run_font(r, "微软雅黑", size=11, bold=True if index <= 3 else False)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_proposal() -> None:
    metrics = load_evaluation_metrics()
    coverage = metrics["coverage"]
    policy_tracks = metrics["policy"]["tracks"]
    retrieval = metrics["ai_pipeline"]["retrieval"]["frozen_final_test"]
    questions = metrics["active_questions"]
    pytest_metrics = metrics["pytest"]
    browser_metrics = metrics["browser"]
    doc = Document()
    configure_document(doc, "惠民智导  作品方案")
    cover(doc, "作品方案", "面向社会救助申请前导办的可信 AI 决策辅助系统")
    add_toc(doc, [
        "项目结论与作品简介", "问题洞察与用户价值", "产品方案与人机交互",
        "系统架构与核心技术", "可信 AI 与安全治理", "功能实现与演示闭环",
        "工程质量与量化评测", "创新性与竞品差异", "应用模式与推广计划",
        "风险边界与后续计划", "交付与复现说明", "参考政策与证据来源",
    ])

    add_section_title(doc, "一 项目结论与作品简介")
    paragraph(doc, "惠民智导已经形成可运行、可复现、可解释、可审计的社会救助申请前导办平台。v2.0.0 覆盖低收入人口认定与监测、最低生活保障、低保边缘家庭、刚性支出困难家庭、临时救助和特困人员救助供养六类事项。系统不替代主管部门审批；它把用户表达整理成待核对事实，把政策条款编译成版本化规则，再用三态逻辑输出“优先核对、还需确认、当前未匹配”，每项结论均可回到政策原文。")
    paragraph(doc, "当前版本已经形成端到端产品闭环：居民任务可以完成，真实本地模型受到规则和证据约束，政策来源可追溯，冻结评测可以复现，候选规则有安全发布边界。项目仍不能宣称已被政府采用或替代人工审批；外部专家盲评、真实用户试点和正式部署评审是下一阶段必须补充的证据。")
    add_section_title(doc, "1.1 作品简介")
    paragraph(doc, "面向社会救助政策分散、群众难理解、材料易缺漏和大模型可能越权判断的问题，惠民智导构建“自然语言辅助录入—三态规则预筛—有据政策问答—本地 OCR 材料预检—Policy2Rule 候选规则治理”闭环。Qwen 本地模型只负责理解与证据化整理，资格方向由版本化规则引擎计算；信息不足保持 UNKNOWN，引用、数值或边界校验失败自动退回逐字证据。平台覆盖六类救助事项，具备本地运行、最小数据、可删除、可审计和可复现评测能力，服务居民申请准备、基层咨询与规则运维。")
    add_picture(doc, "01_服务首页_v2.png", "图 1  居民首页以任务语言回答“先做什么”，避免首次使用者面对复杂功能菜单")

    add_section_title(doc, "二 问题洞察与用户价值")
    add_section_title(doc, "2.1 现有流程的四个断点", 2)
    add_table(doc, ["断点", "用户后果", "平台应对"], [
        ["政策文件分散且有时效", "难判断当前适用版本，容易引用过期条款", "查询日过滤、来源登记、原文指纹和有效期校验"],
        ["用户不知道先说什么", "首次使用看不懂页面，填表放弃率高", "任务化首页、四步向导、可选辅助录入和结果页直达补项"],
        ["生成式模型可能猜测", "把未知当满足、编造数字或越权给资格结论", "AI 与判断解耦，Schema 校验、引用白名单和安全回退"],
        ["规则变化难安全上线", "更新政策可能误伤历史判断", "Policy2Rule 仅生成候选，差异、影响分析和人工发布门"],
    ], [3.0, 5.3, 7.4])
    add_section_title(doc, "2.2 三类核心用户", 2)
    bullets(doc, [
        "居民：匿名完成申请前预筛，理解为什么、还缺什么、下一步去哪里，不要求填写姓名或身份证号。",
        "基层工作人员：在专业工作台并排核对六类事项，保留未知项，减少口径漂移和重复解释。",
        "规则管理员：核对政策来源、候选条件、版本差异、合成案例影响和发布阻断项。",
    ])
    add_section_title(doc, "2.3 可验证的价值", 2)
    add_table(doc, ["对象", "直接价值", "可验证证据"], [
        ["居民", "降低理解和准备成本", "四步向导、补项定位、材料包预检、逐条原文"],
        ["基层", "统一咨询口径", "同一规则快照、三态判断、匿名咨询摘要"],
        ["管理者", "降低规则更新风险", "20 个官方来源、33 条证据、候选规则禁止自动发布"],
        ["赛事评审", "结果可复现", "冻结集、模型摘要、源码指纹、报告版本绑定"],
    ], [2.4, 5.2, 8.1])

    add_section_title(doc, "三 产品方案与人机交互")
    add_section_title(doc, "3.1 信息架构")
    paragraph(doc, "居民侧只保留五个入口：首页、申请前预筛、材料检查、政策问答和可信说明。专业工作台与居民服务分离，避免管理概念干扰普通用户。首页先呈现“先弄清楚方向，再准备材料”的任务目标，再按“不知道能申请什么、已经准备材料、有具体政策问题”三种情境分流。")
    add_section_title(doc, "3.2 首次使用的引导策略", 2)
    bullets(doc, [
        "每一步解释“为什么要问”，未知可留空，不用猜。",
        "特困供养采用先询问对象线索、再展开三项法定条件的渐进披露，避免所有用户都看到复杂字段。",
        "结果先显示结论层级，再显示计算摘要、最小补项、行动建议和政策原文。",
        "提供大字模式、键盘操作、结果朗读与打印；自动化 WCAG A/AA 扫描无违规项。",
    ])
    add_picture(doc, "02_特困供养逐项核对_v2.png", "图 2  特困供养按对象线索渐进展开无劳动能力、无生活来源和法定支持三项核对")
    add_picture(doc, "03_特困供养可信结果_v2.png", "图 3  结果页把结论、未知项、优先事项和下一步行动分层呈现")

    add_section_title(doc, "四 系统架构与核心技术")
    add_section_title(doc, "4.1 受约束的混合架构", 2)
    add_table(doc, ["层", "主要组件", "职责与边界"], [
        ["交互层", "响应式单页界面、四步向导、专业工作台", "收集匿名结构化事实，提供无障碍交互和可见状态"],
        ["API 层", "FastAPI、Pydantic Schema、鉴权和审计", "验证输入、限制高风险操作、提供一致错误模型"],
        ["AI 协作层", "Qwen3 4B、Qwen3 Embedding、本地 OCR", "抽取候选事实、混合检索和证据化整理，不判断资格"],
        ["决策层", "版本化规则注册表、三态规则引擎、关系编排", "计算六类事项，UNKNOWN 不被折叠为通过或拒绝"],
        ["证据与治理层", "政策目录、逐字证据、Policy2Rule、评测报告", "追溯来源、时态校验、候选发布门和版本绑定"],
    ], [2.2, 5.4, 8.1])
    add_section_title(doc, "4.2 核心处理链", 2)
    add_table(doc, ["阶段", "输入", "处理", "输出"], [
        ["1", "匿名家庭情况", "字段校验、最小问题策略", "结构化事实与 UNKNOWN"],
        ["2", "查询日期与地区", "选择有效规则快照", "适用参数和政策版本"],
        ["3", "事实与规则", "Kleene 三态求值与事项关系编排", "优先核对、还需确认、当前未匹配"],
        ["4", "问题与证据库", "混合召回、模型整理、四类校验", "有引用的回答或逐字回退"],
        ["5", "政策原文", "确定性抽取、LLM 抽取、受控融合", "待人工复核候选规则"],
    ], [1.3, 3.8, 6.2, 4.4])
    add_section_title(doc, "4.3 三态规则引擎", 2)
    paragraph(doc, "传统布尔逻辑会把缺失字段误当作 false，进而把“暂时不知道”误写成“不符合”。本项目对每一条件保留 TRUE、FALSE、UNKNOWN；事项结果由条件组合、政策有效期和事项依赖共同决定。输出不仅包含状态，还包含条件代码、解释、补项提示、材料建议、下一步和证据编号。")
    add_section_title(doc, "4.4 Policy2Rule 候选规则编译", 2)
    paragraph(doc, "政策条款进入确定性抽取、LLM 抽取和受控融合三轨。系统要求每个条件和参数必须锚定原文，无法解析的表达保留 unknown；随后执行 JSON Schema、时间区间、重复条件、参数范围、来源指纹、合成案例影响与版本差异检查。候选规则始终带“禁止自动发布”标记，只有人工复核后才可能进入正式注册表。")
    add_picture(doc, "07_Policy2Rule工作台_v2.png", "图 4  Policy2Rule 展示候选条件、原文锚点、版本差异、影响分析和发布阻断")

    add_section_title(doc, "五 可信 AI 与安全治理")
    add_section_title(doc, "5.1 AI 负责理解 规则负责判断", 2)
    paragraph(doc, "模型允许完成四类工作：自然语言字段候选抽取、补充情况中的遗漏与冲突发现、向量检索和证据化表达。模型禁止完成三类工作：直接宣布获得资格、把缺失信息补成符合、绕过规则和人工发布门。所有模型结果先经过结构化校验，用户对候选事实逐项采用或拒绝。")
    add_section_title(doc, "5.2 有据问答的双层护栏", 2)
    bullets(doc, [
        "生成前：领域门、隐私模式、提示注入检测、查询日有效性过滤和混合检索。",
        "生成后：引用必须属于召回白名单；逐字原文必须与证据库一致；数字必须能回到证据；不得产生审批式确定结论。",
        "任何校验失败都退回经核验的逐字证据，宁可降低语言流畅度，也不牺牲来源可核查性。",
    ])
    add_picture(doc, "04_有据政策问答_v2.png", "图 5  真实本地模型生成的特困供养答复，每条结论绑定证据编号")
    add_section_title(doc, "5.3 数据与隐私", 2)
    add_table(doc, ["数据类型", "默认处理", "控制措施"], [
        ["自然语言家庭描述", "仅本次请求使用，不持久化", "敏感字段拦截、最小采集、候选确认"],
        ["上传材料", "本地 OCR，默认不保存原件", "大小页数限制、哈希记录、文件类型校验"],
        ["预筛结构化结果", "用户选择后才保存，最长 30 天", "删除凭证、管理员统计不暴露明细"],
        ["管理接口", "默认拒绝匿名访问", "令牌鉴权、请求审计、发布门"],
    ], [3.2, 5.4, 7.5])
    add_picture(doc, "06_可信机制说明_v2.png", "图 6  产品内公开展示模型、规则、来源、证据与就绪度边界")

    page_break(doc)
    add_section_title(doc, "六 功能实现与演示闭环")
    add_table(doc, ["模块", "关键功能", "演示判据"], [
        ["申请前预筛", "六类事项、四步向导、三态结果、事项关系", "同一页面出现结论、计算依据、未知项和下一步"],
        ["辅助录入", "自然语言转候选事实、原句锚定、冲突提示", "候选不自动写表，用户逐项确认"],
        ["有据问答", "时态过滤、混合检索、模型整理、安全回退", "每条结论有证据编号和逐字原文"],
        ["材料预检", "图片/PDF 质量、OCR、类别和缺项提示", "显示质量指标、识别摘要和人工核对边界"],
        ["专业工作台", "匿名咨询、Policy2Rule、运行治理", "普通用户功能与管理功能隔离"],
    ], [3.0, 7.4, 5.7])
    add_picture(doc, "05_材料智能预检_v2.png", "图 7  本地 OCR 与图像质量预检不保存原始材料")
    add_picture(doc, "08_运行治理面板_v2.png", "图 8  治理面板展示政策来源、证据、规则版本和部署前检查")

    add_section_title(doc, "七 工程质量与量化评测")
    paragraph(doc, "所有指标均限定为项目内冻结集或自动化工程测试，不外推为真实群众满意度或审批准确率。报告与 v2.0.0、规则引擎 0.6.0、模型摘要和源码指纹绑定。")
    add_table(doc, ["评测项", "规模", "当前结果", "边界"], [
        ["Python 自动化测试", f"{pytest_metrics['tests']} 项", "全部通过", "含 API、规则、证据、OCR、安全和回归"],
        ["浏览器端到端", f"{browser_metrics['expected']} 场景", f"{browser_metrics['expected']}/{browser_metrics['expected']} 通过", "含移动端、WCAG A/AA、特困全流程和规则冲突显示"],
        ["代码覆盖率", "全项目", f"行 {coverage['percent_covered']:.2f}%，分支 {coverage['percent_branches_covered']:.2f}%", "质量门检查分支总覆盖不低于 90%"],
        ["意图识别", "300 条冻结测试", "六类事项全覆盖", "项目组合成语句，不代表真实用户分布"],
        ["画像抽取", "300 条 + 特困专题 20 条", "专题精确匹配 20/20", "内部合成并含冲突回归样例"],
        ["政策检索与问答", "138 问 + 专题 20 问", f"BM25 Recall@1 {retrieval['bm25_baseline']['recall_at_1']:.2%}；RRF Recall@1 {retrieval['configured_pipeline']['recall_at_1']:.2%}、Recall@3 {retrieval['configured_pipeline']['recall_at_3']:.2%}", "文件级隔离的内部冻结集"],
        ["真实本地模型专题问答", "20 问", "14 次模型输出通过；6 次安全回退；20/20 闭环通过", "证明护栏有效，不等同外部专家评分"],
        ["Policy2Rule", "120 条冻结政策条款", f"事项识别：确定性 {policy_tracks['deterministic']['benefit_accuracy']:.2%}、纯模型 {policy_tracks['llm_only']['benefit_accuracy']:.2%}、融合 {policy_tracks['controlled_fusion']['benefit_accuracy']:.2%}；条件字段 F1 分别为 {policy_tracks['deterministic']['condition_field']['f1']:.2%}、{policy_tracks['llm_only']['condition_field']['f1']:.2%}、{policy_tracks['controlled_fusion']['condition_field']['f1']:.2%}", "项目组标签；各维度分别报告，禁止自动发布"],
        ["动态提问", f"{questions['cases']} 条合成路径", f"平均 {questions['strategies']['information_gain']['mean_questions']:.2f} 问；相对固定顺序减少 {questions['question_reduction_vs_executed_fixed_order']:.2%}", "仅描述内部合成路径，不代表真人完成率"],
        ["OCR", "100 张，10 种版式", "评测报告可复现", "合成材料，不包含真实个人信息"],
        ["高风险安全", "320 条", "安全报告通过", "覆盖注入、隐私、越权和证据边界"],
    ], [3.2, 3.5, 5.1, 4.3])
    add_section_title(doc, "7.1 可复现机制", 2)
    bullets(doc, [
        "reproduce_phase12.ps1 从数据生成、冻结集、评测、覆盖率到报告签章执行完整复现。",
        "模型名称和 digest 写入 models.lock.json，真实推理烟雾报告记录提示和响应哈希。",
        "release_manifest.json 对代码、测试、评测输入和运行配置计算统一源码指纹。",
        "本次 Windows 本地安全检查使用 Ruff、Mypy、Bandit、定制 AST 安全策略和 secrets 扫描；原生 Semgrep 保留在 Linux CI。",
    ])

    page_break(doc)
    add_section_title(doc, "八 创新性与竞品差异")
    add_table(doc, ["比较维度", "普通政策搜索", "通用大模型问答", "惠民智导"], [
        ["政策时态", "用户自行判断", "可能混用旧知识", "查询日过滤并展示规则快照"],
        ["未知信息", "不处理", "可能补全或猜测", "三态逻辑保留 UNKNOWN 并给最小补项"],
        ["资格边界", "无结构化预筛", "可能越权下结论", "规则引擎判断方向，AI 不决定资格"],
        ["证据链", "有链接但不一定对应结论", "引用可能不可核验", "逐结论证据白名单、原文与数字校验"],
        ["规则更新", "人工改代码", "不可控生成", "三轨候选编译、差异与影响分析、人工发布门"],
        ["数据处理", "依赖网站", "常上传云端", "本地模型与 OCR，最小采集、默认不存原件"],
    ], [3.2, 4.1, 4.3, 5.0])
    add_section_title(doc, "8.1 核心创新点", 2)
    bullets(doc, [
        "把“可信”落实为可执行控制：AI/规则职责隔离、三态逻辑、证据白名单、数值来源校验和安全回退。",
        "把政策更新从一次性代码修改转为可追溯的候选规则治理流程，先影响分析、后人工发布。",
        "把居民、基层和管理员三种任务统一在同一证据与规则底座上，同时隔离界面和权限。",
        "评测资产与源码、模型和规则版本绑定，避免 PPT 指标与实际运行版本脱节。",
    ])

    add_section_title(doc, "九 应用模式与推广计划")
    add_section_title(doc, "9.1 部署方式", 2)
    paragraph(doc, "当前版本适合在高校创新展示、社会救助咨询训练和基层服务原型验证中离线运行。正式政务部署建议采用政务内网或受控专网，接入统一身份、密钥管理、集中日志、灾备和正式政策更新流程。")
    add_section_title(doc, "9.2 分阶段落地", 2)
    add_table(doc, ["阶段", "范围", "验收重点"], [
        ["原型验证", "单机演示与专家评审", "规则、证据、边界和复现"],
        ["封闭试点", "1 至 2 个服务窗口，匿名案例", "可用性、咨询耗时、错误类型和人工纠正"],
        ["受控扩展", "区级多窗口", "统一身份、权限、日志、SLA 和政策更新"],
        ["跨区域适配", "新增地区政策包", "数据隔离、规则差异、地方证据和回归测试"],
    ], [2.5, 5.1, 8.4])
    add_section_title(doc, "9.3 社会与经济价值", 2)
    paragraph(doc, "对居民，价值来自减少无效咨询、重复准备和对政策文本的理解成本；对基层，价值来自统一口径、保留未知和快速定位原文；对管理者，价值来自把规则变更风险前移到候选、差异和合成影响分析阶段。项目当前不虚构节省工时或群众满意度数字，相关收益必须在真实试点中测量。")

    add_section_title(doc, "十 风险边界与后续计划")
    add_table(doc, ["剩余风险", "当前控制", "下一步"], [
        ["缺少外部专家双人盲标", "明确标注内部冻结集，不外推", "制定标注指南，开展独立复核并计算一致性"],
        ["缺少真实用户试点", "不宣称政府采用或真实业务准确率", "在授权场景使用匿名案例开展可用性测试"],
        ["地方政策持续变化", "来源指纹、有效期、规则版本和回归集", "建立定期更新责任人和失效提醒"],
        ["正式部署安全要求更高", "本地模型、最小数据、鉴权与审计", "接入 HTTPS、统一身份、密钥托管、集中日志和灾备"],
        ["小模型表达波动", "强 Schema 和生成后校验，失败回退逐字证据", "扩充难例与红队集，持续优化提示和校验器"],
    ], [4.0, 5.8, 6.2])

    add_section_title(doc, "十一 交付与复现说明")
    bullets(doc, [
        "运行包内含源代码、冻结数据、评测报告、锁定依赖、启动和完整复现脚本。",
        "Windows 首次运行执行 setup.ps1；日常演示执行 start.ps1 -DemoMode。",
        "真实本地 AI 需按 models.lock.json 安装并核对模型摘要；无模型时系统仍以确定性规则和逐字证据安全运行。",
        "最终提交目录同时提供作品方案、答辩 PPT、使用说明、佐证材料、演示视频、主证据表、截图和校验清单。",
    ])

    add_section_title(doc, "十二 参考政策与证据来源")
    paragraph(doc, "项目内登记 20 个官方政策来源和 33 条逐字证据。核心来源包括《中华人民共和国社会救助法》、广东省最低生活保障相关制度、广州市社会救助实施办法及广州市政务服务事项办事指南。每条证据保存来源编号、标题、文号、发布日期、有效期、原文链接、摘录、逐字引文和内容指纹，具体清单见主证据表与运行包 app/data/evidence.json。")
    add_table(doc, ["来源层级", "用途", "核验方式"], [
        ["国家", "社会救助制度、特困人员条件与申请主体", "官方网页或公报、逐字引文、来源指纹"],
        ["广东省", "低保、边缘家庭、刚性支出困难家庭口径", "政策文号、有效期和原文定位"],
        ["广州市", "本地办理范围、材料、时限和窗口", "事项清单、窗口记录和查询日过滤"],
    ], [2.5, 6.5, 7.0])

    doc.save(PROPOSAL)


def build_manual() -> None:
    browser_metrics = load_evaluation_metrics()["browser"]
    doc = Document()
    configure_document(doc, "惠民智导  软件使用说明")
    cover(doc, "软件使用说明", "v2.0.0 Windows 本地运行 部署 演示与故障处理指南")
    add_toc(doc, [
        "文档用途与运行边界", "运行包结构与系统要求", "首次安装与启动",
        "居民端操作", "专业工作台操作", "数据与隐私控制", "完整验收与复现",
        "常见故障处理", "比赛演示流程", "人工提交清单",
    ])
    add_section_title(doc, "一 文档用途与运行边界")
    paragraph(doc, "本说明面向参赛演示人员、评审和本地部署人员，覆盖安装、启动、居民端操作、专业工作台、验证和故障处理。惠民智导用于社会救助申请前导办，不是政务审批系统；页面中的“优先核对”“还需确认”“当前未匹配”均为咨询辅助状态。")
    bullets(doc, [
        "不要录入姓名、身份证号、手机号、银行卡号等个人敏感信息。",
        "正式办理以主管部门受理、调查和审批结果为准。",
        "演示材料均为合成或匿名数据，不包含真实个人档案。",
    ])

    add_section_title(doc, "二 运行包结构与系统要求")
    add_table(doc, ["项目", "要求或位置"], [
        ["操作系统", "Windows 10 或 Windows 11 64 位"],
        ["Python", "3.12，setup.ps1 自动创建项目专用 .venv"],
        ["浏览器", "Microsoft Edge 或 Chrome 当前稳定版"],
        ["内存", "确定性模式建议 8GB；本地 Qwen3 4B 建议 16GB"],
        ["启动脚本", "start.ps1；状态 status.ps1；停止 stop.ps1"],
        ["完整复现", "reproduce_phase12.ps1"],
        ["模型锁定", "models.lock.json"],
        ["评测报告", "data/evaluation"],
    ], [4.1, 12.2])
    paragraph(doc, "运行包不包含个人数据。模型权重体积较大，不随代码包分发；模型名称和 SHA-256 digest 已锁定，安装后由就绪度检查核验。")

    add_section_title(doc, "三 首次安装与启动")
    add_section_title(doc, "3.1 解压", 2)
    bullets(doc, [
        "将运行包完整解压到普通目录，不要直接在 ZIP 内运行。",
        "路径可包含中文，但建议避免过深目录和云盘实时同步目录。",
        "在项目根目录打开 PowerShell。",
    ], numbered=True)
    add_section_title(doc, "3.2 安装依赖", 2)
    paragraph(doc, "执行：  .\\setup.ps1")
    paragraph(doc, "脚本创建 .venv 并按 requirements.lock 安装锁定依赖。若 PowerShell 阻止脚本，可仅对当前进程执行 Set-ExecutionPolicy -Scope Process Bypass，然后重试；不建议修改系统级策略。")
    add_section_title(doc, "3.3 启动确定性演示", 2)
    paragraph(doc, "执行：  .\\start.ps1 -DemoMode")
    paragraph(doc, "默认访问 http://127.0.0.1:8000。端口占用时可执行 .\\start.ps1 -DemoMode -Port 8010。确定性模式不依赖大模型，六类规则、证据、OCR 和治理功能可安全演示。")
    add_section_title(doc, "3.4 启用真实本地 AI", 2)
    bullets(doc, [
        "安装并启动 Ollama 兼容服务。",
        "安装 qwen3:4b-instruct 和 qwen3-embedding:0.6b。",
        "核对 models.lock.json 中的 digest；不一致时就绪度不会标记为通过。",
        "执行 .\\start-ai-runtime.ps1 后再启动平台。",
    ], numbered=True)
    add_section_title(doc, "3.5 状态与停止", 2)
    paragraph(doc, "执行 .\\status.ps1 查看应用、模型和端口状态；执行 .\\stop.ps1 停止由项目脚本启动的服务。")
    add_picture(doc, "01_服务首页_v2.png", "图 1  右上角显示规则引擎或本地 AI 就绪状态")

    add_section_title(doc, "四 居民端操作")
    add_section_title(doc, "4.1 申请前预筛", 2)
    bullets(doc, [
        "从首页点击“开始申请前预筛”。",
        "按家庭情况、收入财产、特殊人员与急难情况、核对生成四步填写。",
        "不知道的字段保持“不清楚”或留空，系统不会擅自猜测。",
        "如需核对特困供养，先选择对象线索，再填写无劳动能力、无生活来源和法定义务人情况。",
        "阅读边界提示并勾选同意后生成结果。保存结果为可选项，未勾选时不保存结构化记录。",
    ], numbered=True)
    add_picture(doc, "02_特困供养逐项核对_v2.png", "图 2  第三步对特困供养采用渐进披露")
    add_section_title(doc, "4.2 读取结果", 2)
    paragraph(doc, "先看顶部结论，再看计算摘要和黄色“信息待确认”区。结果卡按“需要优先核对、还需要补充信息、当前未匹配”分组。展开事项可查看每个条件的状态、解释、补项按钮、材料建议、下一步和政策原文。结果页支持朗读和打印。")
    add_picture(doc, "03_特困供养可信结果_v2.png", "图 3  特困供养预筛结果同时显示结论、未知项和行动建议")
    add_section_title(doc, "4.3 辅助录入与补充情况", 2)
    bullets(doc, [
        "辅助录入只生成候选草稿，不会自动提交，也不会改变资格状态。",
        "补充情况会显示原句依据、置信度和与表单的冲突。",
        "用户必须逐项选择采用或不采用；低置信和矛盾内容保持待确认。",
    ])
    add_section_title(doc, "4.4 有据政策问答", 2)
    bullets(doc, [
        "输入一个具体政策问题，并核对查询日期。",
        "查看结论、证据编号、逐字原文、文号、有效期和原文链接。",
        "页面会显示本地模型、逐字证据回退或拦截原因；不要把答复当作审批决定。",
    ], numbered=True)
    add_picture(doc, "04_有据政策问答_v2.png", "图 4  问答结论逐条绑定政策证据")
    add_section_title(doc, "4.5 材料检查", 2)
    bullets(doc, [
        "选择准备咨询的事项，上传图片或 PDF。单次最多 10 个文件，总计不超过 25MB，PDF 不超过 50 页。",
        "勾选本机处理确认后开始检查。",
        "结果显示清晰度、曝光、对比度、边缘清晰度、OCR 摘要、材料类别和可能缺项。",
        "OCR 仅帮助预检，文字和金额必须人工对照原件，系统不判断真伪。",
    ], numbered=True)
    add_picture(doc, "05_材料智能预检_v2.png", "图 5  合成医疗凭据的本地图像质量与 OCR 预检")

    add_section_title(doc, "五 专业工作台操作")
    add_section_title(doc, "5.1 匿名咨询", 2)
    paragraph(doc, "工作人员可并排查看六类事项。输入必须保持匿名；未知项应保留为空，不得为追求明确结论而代填。摘要用于咨询辅助，不能写入正式审批意见。")
    add_section_title(doc, "5.2 Policy2Rule", 2)
    bullets(doc, [
        "粘贴不含个人信息的政策条款，填写来源编号、地区和标题。",
        "事项提示可留空，由系统识别；本地模型辅助为可选。",
        "点击抽取后核对候选事项、参数、条件、原文锚点、版本差异和合成案例影响。",
        "任何阻断项未清除前禁止发布；本演示版不提供自动发布按钮。",
    ], numbered=True)
    add_picture(doc, "07_Policy2Rule工作台_v2.png", "图 6  候选规则与影响分析工作台")
    add_section_title(doc, "5.3 运行治理", 2)
    paragraph(doc, "进入运行治理需提供 ADMIN_API_TOKEN。面板用于查看记录数量、政策来源、原始材料保存策略、留存上限、部署前检查和当前规则版本。请勿在演示录屏或截图中暴露真实生产令牌。")
    add_picture(doc, "08_运行治理面板_v2.png", "图 7  治理面板展示就绪度和规则登记")

    add_section_title(doc, "六 数据与隐私控制")
    add_table(doc, ["操作", "是否保存", "删除或退出"], [
        ["普通预筛且未勾选保存", "不保存结构化结果", "关闭页面即可"],
        ["勾选保存预筛", "保存最小结构化结果，最长 30 天", "使用结果页删除功能和删除凭证"],
        ["自然语言辅助录入", "原文默认不保存", "离开页面即失去本地草稿"],
        ["材料预检", "不保存上传原件", "请求完成后不保留文件"],
        ["管理员统计", "只显示汇总", "按部署策略清理数据库和日志"],
    ], [4.3, 6.3, 5.7])

    add_section_title(doc, "七 完整验收与复现")
    add_section_title(doc, "7.1 一键质量门", 2)
    paragraph(doc, "在项目根目录执行：  .\\scripts\\run_quality_gates.ps1")
    paragraph(doc, "质量门包含格式、静态类型、单元测试、覆盖率、安全策略、密钥扫描和依赖检查。Windows 本机使用定制 AST 安全策略；原生 Semgrep 在 Linux CI 执行。")
    add_section_title(doc, "7.2 阶段二完整复现", 2)
    paragraph(doc, "执行：  .\\reproduce_phase12.ps1")
    bullets(doc, [
        "重建冻结数据集和来源指纹。",
        "重跑意图、画像、检索、问答、Policy2Rule、OCR、安全、消融和压力评测。",
        "重新生成源码绑定清单并为报告写入版本元数据。",
        "执行项目一致性检查。",
    ], numbered=True)
    add_section_title(doc, "7.3 浏览器端到端", 2)
    paragraph(doc, f"安装 Node 依赖后执行 npm run test:e2e。Windows 若测试器代管 Web 服务不能及时退出，可先手动启动应用，再设置 E2E_REUSE_SERVER=true 运行；当前浏览器回归结果为 {browser_metrics['expected']}/{browser_metrics['expected']} 场景通过。")

    add_section_title(doc, "八 常见故障处理")
    add_table(doc, ["现象", "原因", "处理"], [
        ["页面显示服务未就绪", "后端未启动或端口错误", "运行 status.ps1，核对 URL 和端口"],
        ["本地 AI 未连接", "AI_MODE、模型或 digest 不一致", "检查 Ollama 服务、模型列表和 models.lock.json"],
        ["政策问答退回逐字证据", "生成结果未通过引用、数字或边界校验", "这是安全机制；展开证据原文核对"],
        ["材料识别为空", "图片模糊、曝光异常或 OCR 组件未就绪", "重拍清晰图片并查看 readiness"],
        ["管理员接口 401", "令牌缺失或错误", "重新设置 ADMIN_API_TOKEN 并在面板输入"],
        ["脚本被 PowerShell 阻止", "当前进程执行策略限制", "Set-ExecutionPolicy -Scope Process Bypass"],
        ["端口已被占用", "已有服务进程", "使用 stop.ps1 或指定其他 Port"],
    ], [4.1, 5.0, 7.0])

    add_section_title(doc, "九 比赛演示流程")
    add_table(doc, ["时间", "演示内容", "一句话要点"], [
        ["0:00–0:25", "首页与问题", "用户需要的是下一步，不是再读一遍政策"],
        ["0:25–1:35", "特困供养预筛", "三态规则保留未知，三项法定条件逐项可解释"],
        ["1:35–2:15", "有据政策问答", "本地模型整理，引用与数字失败就回退"],
        ["2:15–2:50", "材料预检", "本地 OCR 检查可读性和缺项，不判断真伪"],
        ["2:50–3:35", "Policy2Rule", "政策变更先编译候选、做影响分析，禁止自动发布"],
        ["3:35–4:10", "治理与评测", "版本、来源、模型和报告都可追溯"],
        ["4:10–4:30", "边界与价值", "说明已完成的工程闭环，以及尚未取得的外部验证"],
    ], [2.8, 5.8, 7.5])

    add_section_title(doc, "十 人工提交清单")
    bullets(doc, [
        "将最终交付目录上传到永久有效网盘并填写提取码。",
        "在报名系统填写作品名称“惠民智导可信政策平台”，不超过 20 字。",
        "复制交付说明中的 300 字以内作品简介。",
        "提交作品方案 PDF、答辩 PPT PDF、3 至 5 分钟 MP4 和佐证材料 PDF。",
        "按学校要求补充队员、指导教师和联系方式；这些信息未由系统擅自填写。",
        "如要求真人出镜或团队配音，按演示脚本重新录制；技术画面和讲解结构可直接复用。",
    ])
    doc.save(MANUAL)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_proposal()
    build_manual()
    print(f"DOCS_OK proposal={PROPOSAL.name} manual={MANUAL.name}")


if __name__ == "__main__":
    main()
